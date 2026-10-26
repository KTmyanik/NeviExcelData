"""Resmi kurum Excel'indeki bağlantılardan kuruluş listelerini indirir (BDDK, SPK, KAP, SEDDK, YÖK, VGM, DERBİS).

Girdi dosyasına eklenen her satır ayrı bir çıktı dosyasıdır. Satırdaki bağlantı
kolonlarının hepsi o dosyada birleştirilir. Yeni kolon, hücresinde adres varsa
kendiliğinden okunur.
"""

from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from difflib import SequenceMatcher
from html import unescape
from pathlib import Path
from time import perf_counter

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]

AYAR_DOSYASI = ROOT / "liste_cekici.json"
SON_DOSYASI = ROOT / "liste_cekici_son.json"
YOK_ONBELLEK = ROOT / "data" / "yok_tur_cache.json"
VARSAYILAN_GIRDI = r"C:\Users\mesut\Downloads\Resmi kurum data.xlsx"
VARSAYILAN_CIKTI = r"C:\Users\mesut\Downloads\ResmiKurumData"
UA = "Mozilla/5.0"
_KILIT = threading.Lock()
_IS = {
    "calisiyor": False,
    "asama": "",
    "baslangic": None,
    "bitis": None,
    "sure": 0,
    "hata": "",
    "satirlar": [],
    "toplam_satir": 0,
    "dosya": 0,
}


def fold(text: str) -> str:
    value = str(text or "").replace("İ", "i").replace("I", "ı")
    value = value.casefold()
    return value.translate(str.maketrans({"ı": "i", "ş": "s", "ğ": "g", "ü": "u", "ö": "o", "ç": "c"}))


def clean(value) -> str:
    text = unescape(str(value or ""))
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def sanitize(name: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]', " ", name)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .")
    return cleaned[:80] or "Liste"


def ayarlari_oku() -> dict:
    data = {"girdi": VARSAYILAN_GIRDI, "cikti_klasoru": VARSAYILAN_CIKTI}
    if AYAR_DOSYASI.exists():
        stored = json.loads(AYAR_DOSYASI.read_text(encoding="utf-8"))
        data["girdi"] = str(stored.get("girdi") or data["girdi"])
        data["cikti_klasoru"] = str(stored.get("cikti_klasoru") or data["cikti_klasoru"])
    return data


def ayarlari_yaz(girdi: str, cikti_klasoru: str) -> dict:
    data = {"girdi": girdi.strip(), "cikti_klasoru": cikti_klasoru.strip()}
    AYAR_DOSYASI.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return data


def _hucre_adresi(cell) -> str:
    if cell.hyperlink is not None and cell.hyperlink.target:
        return str(cell.hyperlink.target).strip()
    value = clean(cell.value)
    if value.startswith("http://") or value.startswith("https://"):
        return value
    return ""


def plani_oku(path: Path) -> list[dict]:
    workbook = load_workbook(path)
    try:
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows())
    finally:
        workbook.close()
    if not rows:
        return []
    headers = [clean(cell.value) for cell in rows[0]]
    plan = []
    for row in rows[1:]:
        texts = []
        sources = []
        for cell in row:
            header = headers[cell.column - 1] if cell.column - 1 < len(headers) else ""
            url = _hucre_adresi(cell)
            value = clean(cell.value)
            if url:
                sources.append({
                    "kolon": cell.coordinate,
                    "baslik": header,
                    "metin": value if not value.startswith("http") else header,
                    "url": url,
                })
            elif value:
                texts.append({"baslik": header, "metin": value})
        if not texts and not sources:
            continue
        name = ""
        hints = []
        for item in texts:
            if not name and "musteri" in fold(item["baslik"]):
                name = item["metin"]
            else:
                hints.append(item["metin"])
        if not name and texts:
            name = texts[0]["metin"]
            hints = [item["metin"] for item in texts[1:]]
        plan.append({
            "satir": row[0].row,
            "ad": name or f"Satır {row[0].row}",
            "ipucu": " ".join(hints),
            "kaynaklar": sources,
        })
    counts = Counter(sanitize(item["ad"]) for item in plan)
    used = set()
    for item in plan:
        base = sanitize(item["ad"])
        if counts[base] > 1:
            base = sanitize(f"{base} - {_ayirici(item['ad'], item['ipucu'])}")
        filename = base
        number = 2
        while filename.casefold() in used:
            filename = f"{base} ({number})"
            number += 1
        used.add(filename.casefold())
        item["dosya"] = f"{filename}.xlsx"
        item["islemler"] = [_islem_etiketi(source["url"], item) for source in item["kaynaklar"]]
    return plan


def _ayirici(ad: str, ipucu: str) -> str:
    filtre = _yok_filtre(ad, ipucu)
    if filtre == "VAKIF":
        return "Vakıf"
    if filtre == "DEVLET":
        return "Devlet"
    words = ipucu.split()
    return " ".join(words[:4]) if words else "liste"


def _islem_etiketi(url: str, item: dict) -> str:
    kind = algila(url)
    if kind == "yok":
        filtre = _yok_filtre(item["ad"], item["ipucu"])
        return "YÖK Atlas · vakıf" if filtre == "VAKIF" else "YÖK Atlas · devlet" if filtre == "DEVLET" else "YÖK Atlas · tüm üniversiteler"
    return {
        "bddk": "BDDK kuruluş listesi",
        "kap": "KAP fon listesi",
        "seddk": "SEDDK şirket listesi",
        "spk": "SPK listesi",
        "vgm": "VGM vakıf listesi",
        "dernek": "DERBİS dernek listesi",
    }.get(kind, "Sayfadaki tablo")


def algila(url: str) -> str:
    host = fold(url)
    if "bddk.org.tr" in host and "/kurulus/liste/" in host:
        return "bddk"
    if "kap.org.tr" in host:
        return "kap"
    if "yokatlas.yok.gov.tr" in host or host.rstrip("/").endswith("yokatlas.yok.gov.tr"):
        return "yok"
    if "seddk.gov.tr" in host:
        return "seddk"
    if "spk.gov.tr" in host:
        return "spk"
    if "vgm.gov.tr" in host:
        return "vgm"
    if "siviltoplum.gov.tr" in host or "dernekler.gov.tr" in host:
        return "dernek"
    return "tablo"


_CURL = shutil.which("curl.exe") or shutil.which("curl") or "curl.exe"


# Kurum ağı bazı siteleri (ör. spk.gov.tr) reddetmek yerine askıda bırakıyor. Bağlantı kurulamazsa
# 15 sn'de vazgeçilir, zaman aşımı yeniden denenmez ve o site bu çalıştırmada bir daha denenmez;
# yoksa SPK'ya giden her adres 3 × 70 sn bekletiyordu.
BAGLANTI_SURESI = 15
_YENIDEN_DENE = {6, 7, 18, 35, 52, 56}  # geçici DNS / bağlantı / aktarım hatası
_ULASILAMAYAN: set[str] = set()


def _site(url: str) -> str:
    match = re.match(r"https?://([^/?#]+)", url, flags=re.I)
    return match.group(1).lower() if match else ""


# Kurum ağı internete proxy'yle çıkıyor; hangi adresin hangi proxy'den gideceğini otomatik yapılandırma
# (PAC) dosyası belirliyor. Tarayıcı bunu Windows'tan alıyor, curl almıyor ve doğrudan bağlanmaya çalışıp
# zaman aşımına düşüyordu. Proxy her site için Windows'un kendi PAC yorumlayıcısına (WinHTTP) sorulur;
# proxy kimlik isterse curl Windows oturumuyla (-U : → SSPI) doğrular.
_PROXY_ORTAM = ("HTTPS_PROXY", "https_proxy", "HTTP_PROXY", "http_proxy", "ALL_PROXY", "all_proxy")
_PROXYLER: dict[str, str | None] = {}


def _winhttp_proxy(url: str, pac_adresi: str | None = None) -> str | None:
    """Windows'un bu adres için seçtiği proxy ("10.13.50.100:8080"); doğrudan erişimse None.

    pac_adresi verilmezse kullanıcının Internet ayarlarındaki sabit proxy, PAC adresi ve otomatik algılama kullanılır.
    """
    import ctypes
    from ctypes import wintypes

    # Windows'un ayırdığı metinler ham işaretçi olarak tutulur: okunup GlobalFree ile bırakılmalı.
    class IeAyari(ctypes.Structure):
        _fields_ = [("fAutoDetect", wintypes.BOOL), ("lpszAutoConfigUrl", ctypes.c_void_p),
                    ("lpszProxy", ctypes.c_void_p), ("lpszProxyBypass", ctypes.c_void_p)]

    class OtomatikAyar(ctypes.Structure):
        _fields_ = [("dwFlags", wintypes.DWORD), ("dwAutoDetectFlags", wintypes.DWORD),
                    ("lpszAutoConfigUrl", wintypes.LPCWSTR), ("lpvReserved", ctypes.c_void_p),
                    ("dwReserved", wintypes.DWORD), ("fAutoLogonIfChallenged", wintypes.BOOL)]

    class ProxyBilgisi(ctypes.Structure):
        _fields_ = [("dwAccessType", wintypes.DWORD), ("lpszProxy", ctypes.c_void_p),
                    ("lpszProxyBypass", ctypes.c_void_p)]

    winhttp = ctypes.WinDLL("winhttp", use_last_error=True)
    kernel32 = ctypes.WinDLL("kernel32")
    kernel32.GlobalFree.argtypes = [ctypes.c_void_p]
    winhttp.WinHttpOpen.restype = ctypes.c_void_p
    winhttp.WinHttpOpen.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.LPCWSTR, wintypes.LPCWSTR, wintypes.DWORD]
    winhttp.WinHttpGetProxyForUrl.argtypes = [ctypes.c_void_p, wintypes.LPCWSTR,
                                              ctypes.POINTER(OtomatikAyar), ctypes.POINTER(ProxyBilgisi)]
    winhttp.WinHttpCloseHandle.argtypes = [ctypes.c_void_p]

    def metin(isaretci) -> str | None:
        return ctypes.wstring_at(isaretci) if isaretci else None

    def serbest(*isaretciler):
        for isaretci in isaretciler:
            if isaretci:
                kernel32.GlobalFree(isaretci)

    sabit = None
    otomatik = False
    if pac_adresi is None:
        ie = IeAyari()
        if winhttp.WinHttpGetIEProxyConfigForCurrentUser(ctypes.byref(ie)):
            pac_adresi = metin(ie.lpszAutoConfigUrl)
            sabit = metin(ie.lpszProxy)
            otomatik = bool(ie.fAutoDetect)
            serbest(ie.lpszAutoConfigUrl, ie.lpszProxy, ie.lpszProxyBypass)
    if pac_adresi or otomatik:
        ayar = OtomatikAyar()
        ayar.dwFlags = (2 if pac_adresi else 0) | (1 if otomatik else 0)  # CONFIG_URL | AUTO_DETECT
        ayar.dwAutoDetectFlags = 3 if otomatik else 0  # DHCP | DNS_A
        ayar.lpszAutoConfigUrl = pac_adresi
        ayar.fAutoLogonIfChallenged = True
        oturum = winhttp.WinHttpOpen(UA, 1, None, None, 0)  # WINHTTP_ACCESS_TYPE_NO_PROXY
        if oturum:
            try:
                bilgi = ProxyBilgisi()
                if winhttp.WinHttpGetProxyForUrl(oturum, url, ctypes.byref(ayar), ctypes.byref(bilgi)):
                    proxy = metin(bilgi.lpszProxy) if bilgi.dwAccessType == 3 else None  # NAMED_PROXY
                    serbest(bilgi.lpszProxy, bilgi.lpszProxyBypass)
                    return _proxy_sec(proxy, url)
            finally:
                winhttp.WinHttpCloseHandle(oturum)
    return _proxy_sec(sabit, url)


def _proxy_sec(liste: str | None, url: str) -> str | None:
    """ "a:1;b:2" ya da "http=a:1;https=b:2" biçiminden bu adrese uyan ilk proxy."""
    if not liste:
        return None
    sema = url.split(":", 1)[0].lower()
    adaylar = [p.strip() for p in re.split(r"[;\s]+", liste) if p.strip()]
    for aday in adaylar:
        if "=" in aday:
            tur, deger = aday.split("=", 1)
            if tur.lower() == sema:
                return deger
    return next((a for a in adaylar if "=" not in a), None)


def _proxy_argumanlari(url: str) -> list[str]:
    if sys.platform != "win32" or any(os.environ.get(ad) for ad in _PROXY_ORTAM):
        return []  # ortamda proxy tanımlıysa curl onu kendisi kullanır
    site = _site(url)
    if not site:
        return []
    anahtar = f"{url.split(':', 1)[0].lower()}://{site}"
    if anahtar not in _PROXYLER:
        try:
            _PROXYLER[anahtar] = _winhttp_proxy(f"{anahtar}/")
        except Exception:
            _PROXYLER[anahtar] = None  # proxy bulunamazsa doğrudan denenir
    proxy = _PROXYLER[anahtar]
    if not proxy:
        return []
    return ["-x", proxy if "://" in proxy else f"http://{proxy}", "--proxy-anyauth", "-U", ":"]


def _curl(args: list[str], timeout: int, stdin: bytes | None = None) -> tuple[str, bytes]:
    site = _site(args[-1])
    if site in _ULASILAMAYAN:
        raise RuntimeError(f"{site} adresine bu çalıştırmada bağlanılamadı; atlandı.")
    proxy = _proxy_argumanlari(args[-1])
    last_error = "curl başarısız oldu."
    code = None
    for attempt in range(3):
        try:
            proc = subprocess.run(
                [_CURL, "-S", "--connect-timeout", str(BAGLANTI_SURESI), *proxy, *args],
                input=stdin,
                capture_output=True,
                timeout=timeout + 15,
                check=False,
            )
        except FileNotFoundError as exc:
            raise RuntimeError("curl bulunamadı.") from exc
        except subprocess.TimeoutExpired:
            last_error = "İstek zaman aşımına uğradı."
            break
        if proc.returncode == 0:
            return proc.stdout.decode("utf-8", errors="replace"), proc.stderr
        code = proc.returncode
        detail = proc.stderr.decode("utf-8", errors="replace").strip()
        last_error = detail[:300] or f"curl çıkış kodu {code}"
        if code not in _YENIDEN_DENE:
            break
    baglanti_yok = code in {6, 7} or (code == 28 and "connect" in last_error.lower())
    if site and baglanti_yok:
        _ULASILAMAYAN.add(site)
        last_error = f"{site} adresine bağlanılamadı (ağ veya güvenlik filtresi engelliyor olabilir). {last_error}"
    raise RuntimeError(last_error)


def http_text(url: str, timeout: int = 70) -> str:
    with tempfile.NamedTemporaryFile(delete=False) as handle:
        path = handle.name
    try:
        code, _err = _curl(
            ["-sL", "--max-time", str(timeout), "-A", UA, "-o", path, "-w", "%{http_code}", url],
            timeout,
        )
        data = Path(path).read_bytes()
    finally:
        Path(path).unlink(missing_ok=True)
    text = data.decode("utf-8", errors="replace")
    if "Web Filter Violation" in text or "Erişim Engellendi" in text:
        raise RuntimeError("Sitenin güvenlik filtresi bu adresi engelledi.")
    status = code.strip()
    if status and not status.startswith("2"):
        raise RuntimeError(f"HTTP {status}")
    return text


def http_json(url: str, body: dict | None = None, timeout: int = 40):
    if body is None:
        text = http_text(url, timeout)
    else:
        payload = json.dumps(body, ensure_ascii=False).encode("utf-8")
        raw, _err = _curl(
            [
                "-sL", "--max-time", str(timeout), "-A", UA,
                "-H", "Content-Type: application/json",
                "-H", "Accept: application/json",
                "--data-binary", "@-",
                "-w", "\n%{http_code}",
                url,
            ],
            timeout,
            payload,
        )
        if "\n" in raw:
            text, status = raw.rsplit("\n", 1)
        else:
            text, status = raw, ""
        if "Web Filter Violation" in text or "Erişim Engellendi" in text:
            raise RuntimeError("Sitenin güvenlik filtresi bu adresi engelledi.")
        if status.strip() and not status.strip().startswith("2"):
            raise RuntimeError(f"HTTP {status.strip()}")
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise RuntimeError("Yanıt JSON değil.") from exc


def form_istegi(url: str, jar: str, form: dict | None = None, headers: tuple[str, ...] = (),
                timeout: int = 70) -> str:
    """Çerezli istek: önce sayfa (çerez ve form anahtarı için), sonra aynı oturumla form gönderimi."""
    args = ["-sL", "--max-time", str(timeout), "-A", UA, "-b", jar, "-c", jar, "-w", "\n%{http_code}"]
    for header in headers:
        args += ["-H", header]
    for key, value in (form or {}).items():
        args += ["--data-urlencode", f"{key}={value}"]
    raw, _err = _curl([*args, url], timeout)
    text, _sep, status = raw.rpartition("\n")
    if "Web Filter Violation" in text or "Erişim Engellendi" in text:
        raise RuntimeError("Sitenin güvenlik filtresi bu adresi engelledi.")
    if status.strip() and not status.strip().startswith("2"):
        raise RuntimeError(f"HTTP {status.strip()}")
    return text


def _html_tables(html: str) -> list[tuple[str, list[list[str]]]]:
    document = re.sub(r"<script[\s\S]*?</script>", "", html, flags=re.I)
    parts = re.split(r"(<table[\s\S]*?</table>)", document, flags=re.I)
    tables = []
    for index, part in enumerate(parts):
        if not part.lower().startswith("<table"):
            continue
        previous = clean(parts[index - 1]) if index else ""
        rows = []
        for tr in re.findall(r"<tr[\s\S]*?</tr>", part, flags=re.I):
            cells = [clean(cell) for cell in re.findall(r"<t[dh][^>]*>([\s\S]*?)</t[dh]>", tr, flags=re.I)]
            cells = [cell for cell in cells if cell]
            if cells:
                rows.append(cells)
        if rows:
            tables.append((previous, rows))
    return tables


def _stems(text: str) -> set[str]:
    return {word[:5] for word in fold(text).split() if len(word) >= 4}


def _bddk_items(chunk: str) -> list[dict]:
    items = []
    for li in re.findall(r'<li class="row">([\s\S]*?)</li>', chunk):
        name = re.search(r'class="satirNo">\s*\d+\.\s*</div>([\s\S]*?)</div>', li)
        if not name:
            continue
        web = re.search(r'href="([^"]+)"', li)
        items.append({"ad": clean(name.group(1)), "web": clean(web.group(1)) if web else ""})
    return items


def _fetch_bddk(url: str, ctx: dict) -> tuple[list[dict], list[str]]:
    html = http_text(url)
    buttons = re.findall(r'<button[^>]*class="btn btn-link"[^>]*>([\s\S]*?)</button>', html)
    if not buttons:
        return _bddk_items(html), []
    sections = []
    for button in buttons:
        title = re.sub(r"\(\s*\d+\s*\)", "", clean(button)).strip()
        position = html.find(title)
        if position < 0:
            continue
        match = re.search(r"<ul>([\s\S]*?)</ul>", html[position:])
        if not match:
            continue
        items = _bddk_items(match.group(1))
        if items:
            sections.append((title, items))
    if not sections:
        return _bddk_items(html), []
    blob = f"{ctx['ad']} {ctx['ipucu']}"
    scored = [(len(_stems(blob) & _stems(title)), title, items) for title, items in sections]
    best = max(score for score, _title, _items in scored)
    chosen = [item for item in scored if item[0] == best]
    if best >= 1 and len(chosen) == 1:
        _score, title, items = chosen[0]
        if "mevduat" in fold(title):
            return _sermaye_ekle(items)
        return items, []
    records = []
    for title, items in sections:
        for item in items:
            records.append({**item, "bolum": title})
    note = "Sayfada birden fazla bölüm var; bölüm ayrımıyla hepsi alındı."
    return records, [note]


# BDDK mevduat bankalarını kamu / özel / yabancı diye ayırmaz; Nevi'de bunlar ayrı tiplerdir (25, 26, 27).
# Ayrım Türkiye Bankalar Birliği'nin üye bankalar sayfasındaki sermaye gruplarından alınır.
TBB_ADRESI = "https://www.tbb.org.tr/banka-ve-sektor-bilgileri/banka-bilgileri/bankalarimiz"
_SERMAYE = (("kamusal", "Kamu"), ("ozel sermayeli", "Özel"), ("yabanci", "Yabancı"), ("tasarruf mevduati", "TMSF"))
_BANKA_EKLERI = {"as", "tas", "tao", "tc", "anonim", "sirketi", "turkiye", "cumhuriyeti", "dijital", "banka"}


def _tbb_bankalari() -> list[dict]:
    """TBB üye bankaları: ad, web adresi ve sayfadaki son grup başlığı (ör. "Kamusal Sermayeli Mevduat Bankaları")."""
    with tempfile.TemporaryDirectory() as klasor:
        jar = str(Path(klasor) / "cerez")
        page = form_istegi(TBB_ADRESI, jar)
        form = page[page.find('id="tbb-custom-blocks-bank-list"'):]
        build = re.search(r'name="form_build_id" value="([^"]+)"', form)
        if not build:
            raise RuntimeError("TBB banka listesi formu bulunamadı.")
        raw = form_istegi(
            f"{TBB_ADRESI}?ajax_form=1&_wrapper_format=drupal_ajax", jar,
            form={
                "date": datetime.now().strftime("%Y-%m-%d"),
                "form_build_id": build.group(1),
                "form_id": "tbb_custom_blocks_bank_list",
                "_triggering_element_name": "op",
                "_triggering_element_value": "Listele",
            },
            headers=("X-Requested-With: XMLHttpRequest", "Accept: application/json"),
        )
    try:
        commands = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise RuntimeError("TBB yanıtı JSON değil.") from exc
    html = next((c["data"] for c in commands if isinstance(c.get("data"), str) and "bank_infos" in c["data"]), "")
    banks = []
    group = ""
    for tr in re.findall(r"<tr[\s\S]*?</tr>", html, flags=re.I):
        cells = re.findall(r"<td[^>]*>([\s\S]*?)</td>", tr, flags=re.I)
        if not cells:
            continue
        if 'class="subheader"' in tr:
            group = clean(cells[0])
            continue
        web = re.search(r'href="(https?://[^"]+)"', tr)
        banks.append({"ad": clean(cells[0]), "web": web.group(1) if web else "", "grup": group})
    if not banks:
        raise RuntimeError("TBB sayfasında banka tablosu bulunamadı.")
    return banks


def _banka_anahtari(ad: str) -> str:
    """Boşluksuz anahtar: "JP MORGAN CHASE BANK" (BDDK) ile "JPMorgan Chase Bank" (TBB) aynı olmalı."""
    words = re.sub(r"[^a-z0-9]+", " ", fold(re.sub(r"\([^)]*\)", " ", ad))).split()
    return "".join(w for w in words if len(w) > 1 and w not in _BANKA_EKLERI)


def _alan_adi(web: str) -> str:
    host = re.sub(r"^https?://", "", fold(web)).split("/")[0]
    return host.removeprefix("www.")


def _tbb_eslesmesi(item: dict, banks: list[dict]) -> dict | None:
    key = _banka_anahtari(item["ad"])
    for bank in banks:
        other = _banka_anahtari(bank["ad"])
        # "INTESA SANPAOLO S.P.A. İTALYA İSTANBUL MERKEZ ŞUBESİ" (BDDK) ~ "Intesa Sanpaolo S.p.A." (TBB)
        if key and other and (key.startswith(other) or other.startswith(key)):
            return bank
    domain = _alan_adi(item.get("web", ""))
    for bank in banks:
        if domain and domain == _alan_adi(bank["web"]):
            return bank
    best = max(banks, key=lambda bank: SequenceMatcher(None, key, _banka_anahtari(bank["ad"])).ratio())
    if SequenceMatcher(None, key, _banka_anahtari(best["ad"])).ratio() >= 0.85:
        return best
    return None


def _sermaye_ekle(items: list[dict]) -> tuple[list[dict], list[str]]:
    try:
        banks = _tbb_bankalari()
    except Exception as exc:
        return items, [f"TBB sermaye grupları alınamadı ({exc}); kamu/özel/yabancı ayrımı yapılmadı."]
    unmatched = []
    for item in items:
        bank = _tbb_eslesmesi(item, banks)
        if bank is None:
            item["sermaye"] = "Belirsiz"
            unmatched.append(item["ad"])
            continue
        item["tbb_grubu"] = bank["grup"]
        item["sermaye"] = next((label for word, label in _SERMAYE if word in fold(bank["grup"])), "Belirsiz")
        if item["sermaye"] == "Belirsiz":
            unmatched.append(item["ad"])
    if unmatched:
        return items, [f"TBB listesinde sermaye grubu bulunamayan bankalar: {', '.join(unmatched)}."]
    return items, []


def _spk_endpoints(html: str) -> list[str]:
    urls = []
    code = re.search(r'kodu:\s*"([^"]+)"', html)
    if code and "KurumListesi/" in html:
        urls.append(f"https://ws.spk.gov.tr/KurumIletisim/api/KurumListesi/{code.group(1)}")
    for url in re.findall(r"https://ws\.spk\.gov\.tr/[A-Za-z0-9_./-]+", html):
        if url.endswith("/"):
            continue
        if any(part in url for part in ("KurulusListe/", "YetkiBelgesi", "KurumListesi/")) and url not in urls:
            urls.append(url)
    return urls


def _json_records(data) -> list[dict]:
    if isinstance(data, dict):
        for key in ("data", "result", "kayitlar", "content"):
            if isinstance(data.get(key), list):
                data = data[key]
                break
    if not isinstance(data, list):
        return []
    records = []
    for item in data:
        if isinstance(item, str):
            if item.strip():
                records.append({"ad": item.strip()})
            continue
        if not isinstance(item, dict):
            continue
        name = item.get("aciklama") or item.get("sirketad") or item.get("unvan") or item.get("ad") or ""
        name = clean(name)
        if not name:
            continue
        record = {"ad": name}
        if item.get("belge"):
            record["belge"] = clean(item.get("belge"))
            record["belge_no"] = clean(item.get("belgeno"))
            record["belge_tarihi"] = clean(item.get("belgetarih"))[:10]
            record["durum"] = clean(item.get("belgedurumu"))
        records.append(record)
    return records


def _table_records(html: str) -> list[dict]:
    tables = _html_tables(html)
    if not tables:
        return []
    _previous, rows = max(tables, key=lambda item: len(item[1]))
    numbered = [row for row in rows if row and row[0][:1].isdigit()]
    if len(numbered) >= 3:
        records = []
        for row in numbered:
            if len(row) >= 3:
                records.append({"ad": row[2], "kurucu": row[1]})
            elif len(row) == 2:
                records.append({"ad": row[1]})
        return records
    if len(rows) < 2:
        return []
    header, body = rows[0], rows[1:]
    records = []
    for row in body:
        if not row:
            continue
        detail = " | ".join(row[1:]) if len(row) > 1 else ""
        record = {"ad": row[0]}
        if detail and not header[0][:1].isdigit():
            labels = [f"{header[index]}: {cell}" for index, cell in enumerate(row) if index > 0 and index < len(header)]
            record["detay"] = " | ".join(labels) if labels else detail
        elif detail:
            record["detay"] = detail
        records.append(record)
    return records


def _fetch_spk(url: str, _ctx: dict) -> tuple[list[dict], list[str]]:
    html = http_text(url)
    warnings = []
    for endpoint in _spk_endpoints(html):
        try:
            records = _json_records(http_json(endpoint))
        except Exception as exc:
            warnings.append(f"{endpoint}: {exc}")
            continue
        if records:
            return records, []
    records = _table_records(html)
    if records:
        return records, []
    warnings.append("Sayfada liste bulunamadı.")
    return [], warnings


def _kap_funds(html: str, fund_type: str) -> list[dict]:
    pattern = re.compile(
        r'\\?"fundCode\\?":\\?"([^"\\]*)\\?",\\?"fundName\\?":\\?"([^"\\]*)\\?",\\?"fundType\\?":\\?"([^"\\]*)\\?"'
    )
    seen = set()
    records = []
    for code, name, kind in pattern.findall(html):
        if fund_type and kind != fund_type:
            continue
        name = clean(name)
        key = (code.strip(), name)
        if not name or key in seen:
            continue
        seen.add(key)
        records.append({"ad": name, "fon_kodu": code.strip()})
    return records


def _fetch_kap(url: str, _ctx: dict) -> tuple[list[dict], list[str]]:
    html = http_text(url, timeout=90)
    match = re.search(r"/YatirimFonlari/([A-Za-z0-9]+)", url, flags=re.I)
    fund_type = match.group(1).upper() if match else ""
    records = _kap_funds(html, fund_type)
    if records:
        return records, []
    return [], ["Fon listesi bulunamadı."]


def _yok_filtre(ad: str, ipucu: str) -> str:
    blob = fold(f"{ad} {ipucu}")
    vakif = "vakif" in blob
    devlet = "devlet" in blob
    if vakif and not devlet:
        return "VAKIF"
    if devlet and not vakif:
        return "DEVLET"
    return ""


def _yok_cache() -> dict:
    if not YOK_ONBELLEK.exists():
        return {}
    try:
        return json.loads(YOK_ONBELLEK.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def _yok_turleri(universities: list[dict]) -> dict[str, str]:
    cache = _yok_cache()
    missing = [item for item in universities if str(item.get("universiteId")) not in cache]

    def one(item):
        uid = item["universiteId"]
        try:
            data = http_json(
                "https://yokatlas.yok.gov.tr/api/tercih-kilavuz/search",
                {"filters": {"universiteId": [uid]}, "page": 0, "size": 1, "sortBy": "basariSirasi", "direction": "ASC"},
            )
            content = data.get("content") or []
            return str(uid), clean(content[0].get("universiteTuru")) if content else ""
        except Exception:
            return str(uid), ""

    if missing:
        with ThreadPoolExecutor(max_workers=8) as pool:
            futures = [pool.submit(one, item) for item in missing]
            for future in as_completed(futures):
                uid, kind = future.result()
                if kind:
                    cache[uid] = kind
        YOK_ONBELLEK.parent.mkdir(parents=True, exist_ok=True)
        YOK_ONBELLEK.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    return cache


def _fetch_yok(_url: str, ctx: dict) -> tuple[list[dict], list[str]]:
    universities = http_json("https://yokatlas.yok.gov.tr/api/tercih-kilavuz/universiteler")
    if not isinstance(universities, list):
        return [], ["Üniversite listesi okunamadı."]
    kinds = _yok_turleri(universities)
    wanted = _yok_filtre(ctx["ad"], ctx["ipucu"])
    records = []
    unknown = 0
    for item in universities:
        name = clean(item.get("universiteAdi"))
        kind = kinds.get(str(item.get("universiteId")), "")
        if not name:
            continue
        if not kind:
            unknown += 1
        if wanted and kind != wanted:
            continue
        record = {"ad": name}
        if not wanted:
            record["tur"] = kind or "Belirsiz"
        records.append(record)
    warnings = []
    if unknown:
        warnings.append(f"{unknown} üniversitenin türü YÖK kaydında boş.")
    if not records:
        warnings.append("Filtreye uyan üniversite bulunamadı.")
    return records, warnings


def _heading(previous: str) -> str:
    tail = clean(previous)
    index = tail.rfind("Şirketleri")
    if index < 0:
        return ""
    chunk = tail[max(0, index - 50):index]
    earlier = chunk.rfind("Şirketleri")
    if earlier >= 0:
        chunk = chunk[earlier + len("Şirketleri"):]
    return clean(chunk + "Şirketleri")


def _fetch_seddk(url: str, ctx: dict) -> tuple[list[dict], list[str]]:
    html = http_text(url)
    wants = []
    blob = fold(f"{ctx['ad']} {ctx['ipucu']}")
    if "sigorta" in blob:
        wants.append("sigorta")
    if "emekli" in blob:
        wants.append("emekli")
    if "reasurans" in blob:
        wants.append("reasurans")
    records = []
    for previous, rows in _html_tables(html):
        heading = _heading(previous)
        folded = fold(heading)
        if not folded.endswith("sirketleri") or "ruhsat" in folded:
            continue
        if wants and not any(word in folded for word in wants):
            continue
        for row in rows:
            records.append({
                "ad": row[0],
                "web": row[-1] if len(row) > 1 else "",
                "tur": heading,
            })
    if records:
        return records, []
    return [], ["İstenen şirket tablosu bulunamadı."]


VGM_SORGU = "https://www.vgm.gov.tr/Foundation/FoundationFilterSearch?lang=tr"
_VGM_SAYFA = 1000  # sunucu bundan büyük sayfa boyutunu reddediyor (7000 → yönlendirme)
_VGM_KATEGORI = {2: "Cemaat", 3: "Mülhak", 4: "Yabancı", 5: "Yeni"}


def _fetch_vgm(url: str, _ctx: dict) -> tuple[list[dict], list[str]]:
    """Vakıflar Genel Müdürlüğü vakıf sorgulaması: sayfa tabloyu JavaScript ile doldurur, veri JSON servistedir."""
    records = []
    with tempfile.TemporaryDirectory() as klasor:
        jar = str(Path(klasor) / "cerez")
        token = re.search(r'name="__RequestVerificationToken" type="hidden" value="([^"]+)"', form_istegi(url, jar))
        if not token:
            raise RuntimeError("Vakıf sorgulama formu bulunamadı.")
        page = 1
        total = None
        while total is None or len(records) < total:
            raw = form_istegi(VGM_SORGU, jar, form={
                "__RequestVerificationToken": token.group(1), "SearchKeyword": "",
                "Page": page, "MaxResultCount": _VGM_SAYFA,
            }, timeout=120)
            try:
                data = json.loads(raw)
            except json.JSONDecodeError as exc:
                raise RuntimeError(f"Vakıf listesi {page}. sayfa JSON değil.") from exc
            items = data.get("Items") or []
            total = int(data.get("TotalCount") or 0)
            for item in items:
                name = clean((item.get("FT_Name_List") or {}).get("TranslatedValue"))
                if name:
                    kategori = item.get("FK_FoundationCategory_Id")
                    records.append({
                        "ad": name,
                        "tur": _VGM_KATEGORI.get(kategori, f"Kategori {kategori}"),
                        "detay": clean(item.get("CityAndState")),
                    })
            if not items:
                break
            page += 1
    if not records:
        return [], ["Vakıf listesi boş geldi."]
    if total and len(records) < total:
        return records, [f"Sitede {total} vakıf görünüyor, {len(records)} tanesi alınabildi."]
    return records, []


DERBIS_SORGU = "https://derbis.dernekler.gov.tr/dernek/api/EdrKurum/GetIlFaaliyetGoreDernekAnonymous"


def _fetch_dernek(_url: str, _ctx: dict) -> tuple[list[dict], list[str]]:
    """İçişleri Bakanlığı DERBİS: "İllere ve Faaliyet Alanlarına Göre Dernekler" sorgusu, tüm iller tek yanıtta."""
    data = http_json(DERBIS_SORGU, {
        "loadOptions": {"filter": {"searchOperation": "contains", "searchValue": None, "userData": {}}},
        "selectedIl": -1,
        "selectedIlce": -1,
    }, timeout=300)
    items = ((data or {}).get("loadResult") or {}).get("data") or []
    records = []
    for item in items:
        name = clean(item.get("kurumAdi"))
        folded = fold(name)
        # Sistemdeki deneme kayıtları ("DERBİS TEST DERNEĞİ ANTALYA ŞUBESİ")
        if not name or ("derbis" in folded and ("test" in folded or "deneme" in folded)):
            continue
        records.append({"ad": name, "tur": clean(item.get("nevi")), "kutuk_no": clean(item.get("kutukNo"))})
    if not records:
        return [], ["Dernek listesi boş geldi."]
    return records, []


def _fetch_tablo(url: str, _ctx: dict) -> tuple[list[dict], list[str]]:
    records = _table_records(http_text(url))
    if records:
        return records, []
    return [], ["Sayfada okunabilir tablo bulunamadı."]


_FETCH = {
    "bddk": _fetch_bddk,
    "spk": _fetch_spk,
    "kap": _fetch_kap,
    "yok": _fetch_yok,
    "seddk": _fetch_seddk,
    "vgm": _fetch_vgm,
    "dernek": _fetch_dernek,
    "tablo": _fetch_tablo,
}


def _etiket(source: dict) -> str:
    parts = []
    for value in (source.get("baslik"), source.get("metin")):
        text = clean(value)
        if text and text not in parts and not text.startswith("http"):
            parts.append(text)
    return " · ".join(parts) or source["url"]


_EK_ALANLAR = (
    ("web", "Web Adresi"),
    ("tur", "Tür"),
    ("sermaye", "Sermaye"),
    ("tbb_grubu", "TBB Grubu"),
    ("kutuk_no", "Kütük No"),
    ("fon_kodu", "Fon Kodu"),
    ("kurucu", "Kurucu"),
    ("belge", "Belge"),
    ("belge_no", "Belge No"),
    ("belge_tarihi", "Belge Tarihi"),
    ("durum", "Durum"),
    ("bolum", "Bölüm"),
    ("detay", "Detay"),
)


def _yaz(path: Path, records: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    used = [item for item in _EK_ALANLAR if any(clean(record.get(item[0])) for record in records)]
    headers = ["Sıra", "Kaynak", "Adres", "Ad", *[label for _key, label in used]]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Liste"
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(1, column, header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for index, record in enumerate(records, 1):
        values = [index, record.get("kaynak", ""), record.get("adres", ""), record.get("ad", "")]
        values.extend(clean(record.get(key)) for key, _label in used)
        for column, value in enumerate(values, 1):
            sheet.cell(index + 1, column, value)
    for column, header in enumerate(headers, 1):
        width = min(70, max(len(header) + 2, 14))
        sheet.column_dimensions[get_column_letter(column)].width = width
    last = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A1:{last}{len(records) + 1}"
    sheet.freeze_panes = "A2"
    gecici = path.with_name(f".{path.name}.yaziliyor")
    workbook.save(gecici)
    os.replace(gecici, path)


def satiri_cek(item: dict, cikti: Path) -> dict:
    started = perf_counter()
    records = []
    warnings = []
    sources = []
    if not item["kaynaklar"]:
        warnings.append("Bu satırda adres yok.")
    for source in item["kaynaklar"]:
        source_started = perf_counter()
        try:
            found, notes = _FETCH[algila(source["url"])](source["url"], item)
            for record in found:
                record["kaynak"] = _etiket(source)
                record["adres"] = source["url"]
            records.extend(found)
            durum = "uyari" if notes else "tamam"
            warnings.extend(notes)
            error = ""
        except Exception as exc:
            found = []
            durum = "hata"
            error = str(exc)
            warnings.append(f"{_etiket(source)}: {exc}")
        sources.append({
            "etiket": _etiket(source),
            "url": source["url"],
            "satir": len(found),
            "sure": round(perf_counter() - source_started, 1),
            "durum": durum,
            "hata": error,
        })
    target = cikti / item["dosya"]
    if records:
        try:
            _yaz(target, records)
        except PermissionError:
            warnings.append(f"{item['dosya']} açık olduğu için kaydedilemedi.")
            target = None
    elif target.exists():
        target = None
    durum = "hata" if not records else "uyari" if warnings else "tamam"
    return {
        "satir": item["satir"],
        "ad": item["ad"],
        "dosya": item["dosya"] if records and target else "",
        "yol": str(target) if records and target else "",
        "satir_sayisi": len(records),
        "sure": round(perf_counter() - started, 1),
        "durum": durum,
        "uyarilar": warnings,
        "kaynaklar": sources,
    }


def calistir(girdi: str, cikti_klasoru: str, ilerleme=None) -> dict:
    source = Path(girdi.strip())
    target = Path(cikti_klasoru.strip())
    if not source.is_file():
        raise FileNotFoundError(f"Girdi dosyası bulunamadı: {source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Girdi dosyası .xlsx olmalı.")
    started = perf_counter()
    plan = plani_oku(source)
    _ULASILAMAYAN.clear()  # ağ değişmiş olabilir; her çalıştırma siteleri ve proxy'yi yeniden dener
    _PROXYLER.clear()
    results = []
    for index, item in enumerate(plan, 1):
        if ilerleme:
            ilerleme(f"{index}/{len(plan)} · {item['ad']}", results)
        results.append(satiri_cek(item, target))
        if ilerleme:
            ilerleme(f"{index}/{len(plan)} · {item['ad']} bitti", results)
    return {
        "sure": round(perf_counter() - started, 1),
        "dosya": sum(1 for item in results if item["dosya"]),
        "toplam_satir": sum(item["satir_sayisi"] for item in results),
        "satirlar": results,
        "hata": "",
    }


def is_durumu() -> dict:
    with _KILIT:
        if _IS["baslangic"] is None and SON_DOSYASI.exists():
            stored = json.loads(SON_DOSYASI.read_text(encoding="utf-8"))
            _IS.update(stored)
            _IS["calisiyor"] = False
            _IS["baslangic"] = 0
        data = json.loads(json.dumps(_IS))
    if data.get("calisiyor") and data.get("baslangic"):
        data["sure"] = round(perf_counter() - data["baslangic"], 1)
    return data


def isi_baslat(girdi: str, cikti_klasoru: str) -> dict:
    with _KILIT:
        if _IS["calisiyor"]:
            raise RuntimeError("Bir indirme zaten çalışıyor.")
        _IS.update({
            "calisiyor": True,
            "asama": "Excel okunuyor",
            "baslangic": perf_counter(),
            "bitis": None,
            "sure": 0,
            "hata": "",
            "satirlar": [],
            "toplam_satir": 0,
            "dosya": 0,
        })

    def ilerleme(asama, satirlar):
        with _KILIT:
            _IS["asama"] = asama
            _IS["satirlar"] = list(satirlar)
            _IS["toplam_satir"] = sum(item["satir_sayisi"] for item in satirlar)
            _IS["dosya"] = sum(1 for item in satirlar if item["dosya"])

    def run():
        try:
            result = calistir(girdi, cikti_klasoru, ilerleme)
            with _KILIT:
                _IS.update(result)
        except Exception as exc:
            with _KILIT:
                _IS["hata"] = str(exc)
                _IS["asama"] = "Hata"
        finally:
            with _KILIT:
                _IS["calisiyor"] = False
                _IS["bitis"] = perf_counter()
                if _IS.get("baslangic"):
                    _IS["sure"] = round(_IS["bitis"] - _IS["baslangic"], 1)
                payload = {key: value for key, value in _IS.items() if key != "baslangic"}
            SON_DOSYASI.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    threading.Thread(target=run, daemon=True).start()
    return is_durumu()
